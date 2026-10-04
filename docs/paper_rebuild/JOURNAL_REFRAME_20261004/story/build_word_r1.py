from pathlib import Path
import re,json,hashlib,zipfile,datetime,copy
from docx import Document
from docx.shared import Inches,Pt,RGBColor,Cm
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
D=Path(__file__).resolve().parent
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
src=D/'manuscript_restructured_r1.md';before=sha(src);t=src.read_text();doc=Document();sec=doc.sections[0]
sec.page_width=Cm(21);sec.page_height=Cm(29.7);sec.top_margin=Cm(2);sec.bottom_margin=Cm(2);sec.left_margin=Cm(2.2);sec.right_margin=Cm(2.2)
for n in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption']:
 s=doc.styles[n];s.font.name='Times New Roman';s.font.color.rgb=RGBColor(0,0,0);s.font.size=Pt(11.5)
 s.paragraph_format.space_after=Pt(5);s.paragraph_format.line_spacing=1.08
 for e in list(s.element.xpath('.//w:pBdr')):e.getparent().remove(e)
doc.styles['Title'].font.size=Pt(19);doc.styles['Title'].font.bold=True;doc.styles['Title'].paragraph_format.space_after=Pt(12)
doc.styles['Subtitle'].font.size=Pt(10);doc.styles['Subtitle'].paragraph_format.space_after=Pt(16)
for n,sz in [('Heading 1',13),('Heading 2',11.5)]:
 doc.styles[n].font.size=Pt(sz);doc.styles[n].font.bold=True;doc.styles[n].paragraph_format.space_before=Pt(10);doc.styles[n].paragraph_format.keep_with_next=True
# Native editable OMML builders. No raster equations or raw LaTeX in DOCX.
def el(tag,*children):
 x=OxmlElement('m:'+tag)
 for y in children:x.append(y)
 return x
def mt(s):
 r=el('r');pr=el('rPr');pr.append(el('nor'));r.append(pr);tx=el('t');tx.text=s;r.append(tx);return r
def wrap(tag,items):return el(tag,*items)
def sub(b,s):return el('sSub',wrap('e',b if isinstance(b,list) else [b]),wrap('sub',s if isinstance(s,list) else [mt(s)]))
def sup(b,s):return el('sSup',wrap('e',b if isinstance(b,list) else [b]),wrap('sup',s if isinstance(s,list) else [mt(s)]))
def frac(n,d):return el('f',wrap('num',n if isinstance(n,list) else [mt(n)]),wrap('den',d if isinstance(d,list) else [mt(d)]))
def hat(b):
 x=el('acc');pr=el('accPr');c=el('chr');c.set(qn('m:val'),'̂');pr.append(c);x.append(pr);x.append(wrap('e',[mt(b)]));return x
def scripts(b,low,high):return el('sSubSup',wrap('e',[mt(b)]),wrap('sub',[mt(low)]),wrap('sup',[mt(high)]))
def matrix_transpose(name):return sup(sub(mt('δb'),name),'T')
def mrun(p,items,display=False):
 math=el('oMath',*items)
 if display:
  para=el('oMathPara');para.append(math);p._p.append(para)
 else:p._p.append(math)
 return math
EQ={1:[mt('δx = ['),sup(mt('δp'),'T'),mt(', '),sup(mt('δv'),'T'),mt(', '),sup(mt('δφ'),'T'),mt(', '),matrix_transpose('g'),mt(', '),matrix_transpose('a'),mt(']'),mt('ᵀ'),mt('     (1)')],
 2:[sub(mt('h'),'p'),mt('(x) = '),sub(mt('p'),'IMU'),mt(' + '),scripts('C','b','n'),sub(mt('ℓ'),'1'),mt('     (2)')],
 3:[sub(mt('ψ'),'b'),mt(' = wrap[atan2('),sub(mt('b'),'E'),mt(', '),sub(mt('b'),'N'),mt(') + '),frac('π','2'),mt(']     (3)')],
 4:[sub(mt('r'),'ψ'),mt(' = wrap('),hat('ψ'),mt(' − '),sub(mt('ψ'),'b'),mt(')     (4)')],
 5:[sub(mt('v'),'H'),mt(' = '),sub(mt('Π'),'H'),hat('C'),sub(mt('k'),'HV'),sub(mt('v'),'FLU'),mt(' ; '),hat('C'),mt(' = '),sub(mt('R'),'z'),mt('('),sub(mt('ψ'),'A'),mt(')'),sub(mt('R'),'y'),mt('(−'),sub(mt('θ'),'SDK'),mt(')'),sub(mt('R'),'x'),mt('('),sub(mt('φ'),'SDK'),mt(')M     (5)')],
 6:[mt('R′ = aR ; a = min('),sub(mt('a'),'max'),mt(', max[1, '),sub(mt('a'),'meta'),mt(', '),sub(mt('a'),'innov'),mt('])     (6)')]}
INLINE={'\\ell_1':[sub(mt('ℓ'),'1')],'b=p_2-p_1':[mt('b = '),sub(mt('p'),'2'),mt(' − '),sub(mt('p'),'1')],'\\psi_b':[sub(mt('ψ'),'b')],'M=\\operatorname{diag}(1,-1,-1)':[mt('M = diag(1, −1, −1)')],'\\Pi_H':[sub(mt('Π'),'H')],'\\psi_A':[sub(mt('ψ'),'A')],'k_{\\mathrm{HV}}':[sub(mt('k'),'HV')],'d_A=\\hat y_A-y_R':[sub(mt('d'),'A'),mt(' = '),sub(hat('y'),'A'),mt(' − '),sub(mt('y'),'R')],'d_B=\\hat y_B-y_R':[sub(mt('d'),'B'),mt(' = '),sub(hat('y'),'B'),mt(' − '),sub(mt('y'),'R')],'d_A-d_B':[sub(mt('d'),'A'),mt(' − '),sub(mt('d'),'B')],'y_R':[sub(mt('y'),'R')]}
def para_text(p,s):
 for part in re.split(r'(\$[^$]+\$|\*\*[^*]+\*\*)',s):
  if not part:continue
  if part.startswith('$'):
   v=part[1:-1];assert v in INLINE,v;mrun(p,[copy.deepcopy(x) for x in INLINE[v]])
  elif part.startswith('**'):p.add_run(part[2:-2]).bold=True
  else:p.add_run(part)
 p.paragraph_format.widow_control=True
lines=t.splitlines();i=0;display_count=0;table_count=0
while i<len(lines):
 l=lines[i].strip()
 if not l:i+=1;continue
 if l.startswith('# '):
  doc.add_paragraph(l[2:],style='Title');doc.add_paragraph('Research manuscript draft   Descriptive title provisional',style='Subtitle');i+=1;continue
 if l.startswith('### '):doc.add_paragraph(l[4:],style='Heading 2');i+=1;continue
 if l.startswith('## '):doc.add_paragraph(l[3:],style='Heading 1');i+=1;continue
 if l.startswith('$$'):
  n=int(re.search(r'\\tag\{(\d+)\}',l).group(1));p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_before=Pt(5);p.paragraph_format.space_after=Pt(8);mrun(p,EQ[n],True);display_count+=1;i+=1;continue
 if l.startswith('|'):
  rows=[]
  while i<len(lines) and lines[i].strip().startswith('|'):
   rr=[x.strip() for x in lines[i].strip().strip('|').split('|')]
   if not all(re.fullmatch('[-:]+',x) for x in rr):rows.append(rr)
   i+=1
  tb=doc.add_table(rows=len(rows),cols=len(rows[0]));tb.alignment=WD_TABLE_ALIGNMENT.CENTER;tb.style='Table Grid';tb.autofit=False
  widths=[2.1,1.65,2.3,2.15,2.15,2.15,2.1]
  for k,row in enumerate(rows):
   for z,value in enumerate(row):
    c=tb.cell(k,z);c.width=Cm(widths[z]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;c.text=value
    for p in c.paragraphs:
     p.paragraph_format.space_after=Pt(4);p.paragraph_format.space_before=Pt(4);p.alignment=WD_ALIGN_PARAGRAPH.LEFT if z==0 else WD_ALIGN_PARAGRAPH.CENTER
     for r in p.runs:r.font.size=Pt(9.5);r.font.bold=k==0
    if k==0:
     shading=OxmlElement('w:shd');shading.set(qn('w:fill'),'E5E7EB');c._tc.get_or_add_tcPr().append(shading)
  tr=OxmlElement('w:tblHeader');tb.rows[0]._tr.get_or_add_trPr().append(tr);table_count+=1;continue
 p=doc.add_paragraph();para_text(p,l)
 if l.startswith('**Table '):p.paragraph_format.keep_with_next=True
 i+=1
# Native equations must have all keys represented.
assert display_count==6 and table_count==1
p=doc.add_paragraph();p.paragraph_format.page_break_before=True
doc.add_paragraph('Legends for planned main figures',style='Heading 1')
doc.add_paragraph('The legends describe the eight planned main figures. The text draft does not include newly rendered figures; their retained sources and panel roles are specified in the separate figure plan.')
figs=json.loads((D/'MAIN_FIGURE_PLAN.json').read_text())['figures']
for x in figs:
 p=doc.add_paragraph('Figure '+str(x['figure'])+' '+x['title_en'],style='Heading 2');p.paragraph_format.keep_with_next=True
 doc.add_paragraph(x['legend_en'])
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=footer.add_run('Research draft   ');r.font.size=Pt(9)
f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');footer._p.append(f)
doc.core_properties.title='Status qualified short baseline dual antenna GNSS INS for quadruped navigation';doc.core_properties.subject='New manuscript text draft with editable equations and table';doc.core_properties.author='';doc.core_properties.last_modified_by=''
out=D/'manuscript_restructured_r1.docx';assert not out.exists();doc.save(out)
with zipfile.ZipFile(out) as z:
 xml=z.read('word/document.xml');assert b'\\operatorname' not in xml and b'\\tag' not in xml and b'[REF' not in xml and b'[NEED' not in xml
 assert xml.count(b'<m:oMathPara>')==6
assert sha(src)==before
receipt=dict(status='DOCX_CREATED_RENDER_PENDING',source_md_sha256=before,docx_sha256=sha(out),native_display_equations=6,native_inline_math_objects=xml.count(b'<m:oMath>')-6,editable_tables=1,reference_items=11,planned_figure_legends=8,new_rendered_figures_in_docx=0,scientific_execution=0,builder_sha256=sha(Path(__file__)),source_md_unchanged=True)
(D/'DOCX_R1_AUTHORING_RECEIPT.json').write_text(json.dumps(receipt,ensure_ascii=False,indent=2)+'\n')
print(json.dumps(receipt,ensure_ascii=False))
