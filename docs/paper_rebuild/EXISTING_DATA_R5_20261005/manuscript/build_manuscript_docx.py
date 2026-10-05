from pathlib import Path
import argparse
import re,json,hashlib,zipfile,datetime,copy
from docx import Document
from docx.shared import Inches,Pt,RGBColor,Cm
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
ap=argparse.ArgumentParser();ap.add_argument('--language',choices=['en','zh'],default='en');ap.add_argument('--double-spaced',action='store_true');ap.add_argument('--overwrite',action='store_true');ap.add_argument('--source',type=Path,required=True);ap.add_argument('--output',type=Path,required=True);ap.add_argument('--receipt',type=Path,required=True);args=ap.parse_args()
D=Path(__file__).resolve().parent
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
src=args.source;before=sha(src);t=src.read_text(encoding='utf-8');doc=Document();sec=doc.sections[0]
sec.page_width=Cm(21);sec.page_height=Cm(29.7);sec.top_margin=Cm(2);sec.bottom_margin=Cm(2);sec.left_margin=Cm(2.2);sec.right_margin=Cm(2.2)
for n in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption']:
 s=doc.styles[n];s.font.name='Times New Roman';s.font.color.rgb=RGBColor(0,0,0);s.font.size=Pt(11.5)
 fonts=s.element.get_or_add_rPr().rFonts
 fonts.set(qn('w:eastAsia'),'SimSun')
 for theme in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
  fonts.attrib.pop(qn('w:'+theme),None)
 s.paragraph_format.space_after=Pt(5);s.paragraph_format.line_spacing=2.0 if args.language=='en' else 1.5
 for e in list(s.element.xpath('.//w:pBdr')):e.getparent().remove(e)
doc.styles['Title'].font.size=Pt(19);doc.styles['Title'].font.bold=True;doc.styles['Title'].paragraph_format.space_after=Pt(12)
doc.styles['Subtitle'].font.size=Pt(10);doc.styles['Subtitle'].paragraph_format.space_after=Pt(16)
for n,sz in [('Heading 1',13),('Heading 2',11.5)]:
 doc.styles[n].font.size=Pt(sz);doc.styles[n].font.bold=True;doc.styles[n].paragraph_format.space_before=Pt(10);doc.styles[n].paragraph_format.keep_with_next=True
# Native editable OMML builders. No raster equations or raw LaTeX in DOCX.
def el(tag,*children):
 x=OxmlElement('m:'+tag)
 for y in children:x.append(y)
 return x
def mt(s):
 r=el('r');pr=el('rPr');pr.append(el('nor'));r.append(pr);tx=el('t');tx.text=s;r.append(tx);return r
def wrap(tag,items):return el(tag,*items)
def sub(b,s):return el('sSub',wrap('e',b if isinstance(b,list) else [b]),wrap('sub',s if isinstance(s,list) else [mt(s)]))
def sup(b,s):return el('sSup',wrap('e',b if isinstance(b,list) else [b]),wrap('sup',s if isinstance(s,list) else [mt(s)]))
def frac(n,d):return el('f',wrap('num',n if isinstance(n,list) else [mt(n)]),wrap('den',d if isinstance(d,list) else [mt(d)]))
def hat(b):
 x=el('acc');pr=el('accPr');c=el('chr');c.set(qn('m:val'),'̂');pr.append(c);x.append(pr);x.append(wrap('e',[mt(b)]));return x
def scripts(b,low,high):return el('sSubSup',wrap('e',[mt(b)]),wrap('sub',[mt(low)]),wrap('sup',[mt(high)]))
def matrix_transpose(name):return sup(sub(mt('δb'),name),'T')
def mrun(p,items,display=False):
 math=el('oMath',*items)
 if display:
  para=el('oMathPara');para.append(math);p._p.append(para)
 else:p._p.append(math)
 return math
EQ={1:[mt('δx = ['),sup(mt('δp'),'T'),mt(', '),sup(mt('δv'),'T'),mt(', '),sup(mt('δφ'),'T'),mt(', '),matrix_transpose('g'),mt(', '),matrix_transpose('a'),mt(']'),mt('ᵀ'),mt('     (1)')],
 2:[sub(mt('h'),'p'),mt('(x) = '),sub(mt('p'),'IMU'),mt(' + '),scripts('C','b','n'),sub(mt('ℓ'),'1'),mt('     (2)')],
 3:[sub(mt('ψ'),'b'),mt(' = wrap[atan2('),sub(mt('b'),'E'),mt(', '),sub(mt('b'),'N'),mt(') + '),frac('π','2'),mt(']     (3)')],
 4:[sub(mt('r'),'ψ'),mt(' = wrap('),hat('ψ'),mt(' − '),sub(mt('ψ'),'b'),mt(')     (4)')],
 5:[sub(mt('v'),'H'),mt(' = '),sub(mt('Π'),'H'),hat('C'),sub(mt('k'),'HV'),sub(mt('v'),'FLU'),mt(' ; '),hat('C'),mt(' = '),sub(mt('R'),'z'),mt('('),sub(mt('ψ'),'A'),mt(')'),sub(mt('R'),'y'),mt('(−'),sub(mt('θ'),'SDK'),mt(')'),sub(mt('R'),'x'),mt('('),sub(mt('φ'),'SDK'),mt(')M     (5)')],
 6:[mt('R′ = aR ; a = min('),sub(mt('a'),'max'),mt(', max[1, '),sub(mt('a'),'meta'),mt(', '),sub(mt('a'),'innov'),mt('])     (6)')]}

EQ[7]=[el('eqArr',wrap('e',[sub(mt('U'),'b'),mt(' = '),sub(mt('U'),'22'),mt(' + '),sub(mt('U'),'11'),mt(' − '),sub(mt('U'),'21'),mt(' − '),sub(mt('U'),'12')]),wrap('e',[sup(mt('u'),'2'),mt('('),sub(mt('ψ'),'b'),mt(') ≃ '),sub(mt('j'),'b'),sub(mt('U'),'b'),sup(sub(mt('j'),'b'),'T'),mt(' ; '),sub(mt('j'),'b'),mt(' = [−'),frac([sub(mt('b'),'E')],[sup(mt('r'),'2')]),mt(', '),frac([sub(mt('b'),'N')],[sup(mt('r'),'2')]),mt(', 0]     (7)')]))]
EQ[8]=[sub(mt('t'),'G'),mt(' = a '),sub(mt('t'),'R'),mt(' + c + '),sub(mt('d'),'eff'),mt(' + η     (8)')]
EQ[9]=[sub(mt('U'),'d'),mt(' = '),sub(mt('U'),'m'),mt(' + '),sub(mt('U'),'R'),mt(' − '),sub(mt('U'),'mR'),mt(' − '),sub(mt('U'),'Rm'),mt('     (9)')]

INLINE={'\\ell_1':[sub(mt('ℓ'),'1')],'b=p_2-p_1':[mt('b = '),sub(mt('p'),'2'),mt(' − '),sub(mt('p'),'1')],'\\psi_b':[sub(mt('ψ'),'b')],'M=\\operatorname{diag}(1,-1,-1)':[mt('M = diag(1, −1, −1)')],'\\Pi_H':[sub(mt('Π'),'H')],'\\psi_A':[sub(mt('ψ'),'A')],'k_{\\mathrm{HV}}':[sub(mt('k'),'HV')],'d_A=\\hat y_A-y_R':[sub(mt('d'),'A'),mt(' = '),sub(hat('y'),'A'),mt(' − '),sub(mt('y'),'R')],'d_B=\\hat y_B-y_R':[sub(mt('d'),'B'),mt(' = '),sub(hat('y'),'B'),mt(' − '),sub(mt('y'),'R')],'d_A-d_B':[sub(mt('d'),'A'),mt(' − '),sub(mt('d'),'B')],'y_R':[sub(mt('y'),'R')]}
INLINE.update({'C_b^n':[scripts('C','b','n')],'S_s':[sub(mt('S'),'s')],'d_s':[sub(mt('d'),'s')]})
INLINE['z_s=\\sqrt{r_s^\\mathsf{T}S_s^{-1}r_s/d_s}']=[sub(mt('z'),'s'),mt(' = '),el('rad',el('deg'),wrap('e',[frac([sup(sub(mt('r'),'s'),'T'),sup(sub(mt('S'),'s'),'−1'),sub(mt('r'),'s')],[sub(mt('d'),'s')])]))]

INLINE['[0.03,0.03,-0.30]']=[mt('[0.03, 0.03, −0.30]')]
INLINE['a_{\\mathrm{meta}}']=[sub(mt('a'),'meta')]
INLINE['z_s=\\sqrt{r_s^\\mathsf TS_s^{-1}r_s/d_s}']=INLINE['z_s=\\sqrt{r_s^\\mathsf{T}S_s^{-1}r_s/d_s}']

def para_text(p,s):
 for part in re.split(r'(\$[^$]+\$|\*\*[^*]+\*\*)',s):
  if not part:continue
  if part.startswith('$'):
   v=part[1:-1];assert v in INLINE,v;mrun(p,[copy.deepcopy(x) for x in INLINE[v]])
  elif part.startswith('**'):p.add_run(part[2:-2]).bold=True
  else:p.add_run(part)
 p.paragraph_format.widow_control=True
display_normalized=re.sub(r'\$\$(.*?)\$\$',lambda m:'$$'+re.sub(r'\s+', ' ',m.group(1)).strip()+'$$',t,flags=re.S)
lines=display_normalized.splitlines();i=0;display_count=0;table_count=0
while i<len(lines):
 l=lines[i].strip()
 if not l:i+=1;continue
 if l.startswith('# '):
  doc.add_paragraph(l[2:],style='Title');i+=1;continue
 if l.startswith('### '):doc.add_paragraph(l[4:],style='Heading 2');i+=1;continue
 if l.startswith('## '):doc.add_paragraph(l[3:],style='Heading 1');i+=1;continue
 if l.startswith('$$'):
  n=int(re.search(r'\\tag\{(\d+)\}',l).group(1));p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_before=Pt(5);p.paragraph_format.space_after=Pt(8);mrun(p,EQ[n],True);display_count+=1;i+=1;continue
 if l.startswith('|'):
  rows=[]
  while i<len(lines) and lines[i].strip().startswith('|'):
   rr=[x.strip() for x in lines[i].strip().strip('|').split('|')]
   if not all(re.fullmatch('[-:]+',x) for x in rr):rows.append(rr)
   i+=1
  tb=doc.add_table(rows=len(rows),cols=len(rows[0]));tb.alignment=WD_TABLE_ALIGNMENT.CENTER;tb.style='Table Grid';tb.autofit=False
  widths={0:[8.0,4.3,4.3],1:[2.2,4.7,3.0,3.0,3.7],2:[8.2,2.8,2.8,2.8],3:[6.3,1.6,3.0,3.0,2.7],4:[3.8,2.2,3.5,3.5,3.6]}.get(table_count,[16.6/len(rows[0])]*len(rows[0]))
  for k,row in enumerate(rows):
   no_split=OxmlElement('w:cantSplit');tb.rows[k]._tr.get_or_add_trPr().append(no_split)
   for z,value in enumerate(row):
    c=tb.cell(k,z);c.width=Cm(widths[z]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;c.text=value
    for p in c.paragraphs:
     p.paragraph_format.keep_with_next=k<len(rows)-1;p.paragraph_format.line_spacing=1.0;p.paragraph_format.space_after=Pt(4);p.paragraph_format.space_before=Pt(4);p.alignment=WD_ALIGN_PARAGRAPH.LEFT if z==0 else WD_ALIGN_PARAGRAPH.CENTER
     for r in p.runs:r.font.size=Pt(9.5);r.font.bold=k==0
    if k==0:
     shading=OxmlElement('w:shd');shading.set(qn('w:fill'),'E5E7EB');c._tc.get_or_add_tcPr().append(shading)
  borders=OxmlElement('w:tblBorders')
  for edge in ['top','left','bottom','right','insideH','insideV']:
   border=OxmlElement('w:'+edge);border.set(qn('w:val'),'single');border.set(qn('w:sz'),'4');border.set(qn('w:color'),'D9D9D9');borders.append(border)
  tb._tbl.tblPr.append(borders)
  margins=OxmlElement('w:tblCellMar')
  for edge,value in [('top','70'),('bottom','70'),('left','85'),('right','85')]:
   v=OxmlElement('w:'+edge);v.set(qn('w:w'),value);v.set(qn('w:type'),'dxa');margins.append(v)
  tb._tbl.tblPr.append(margins)
  tr=OxmlElement('w:tblHeader');tb.rows[0]._tr.get_or_add_trPr().append(tr);table_count+=1;continue
 p=doc.add_paragraph();para_text(p,l)
 if l.startswith(('**Table ','**表')):p.paragraph_format.keep_with_next=True
 i+=1
# Native equations must have all keys represented.
expected_equations=len(re.findall(r'\\tag\{\d+\}',t));assert display_count==expected_equations and table_count>=5
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=footer.add_run('');r.font.size=Pt(9)
f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');footer._p.append(f)
doc.core_properties.title=t.splitlines()[0].removeprefix('# ');doc.core_properties.subject='New manuscript text draft with editable equations and table';doc.core_properties.author='';doc.core_properties.last_modified_by=''
out=args.output;out.parent.mkdir(parents=True,exist_ok=True);assert not out.exists() or args.overwrite;doc.save(out)
with zipfile.ZipFile(out) as z:
 xml=z.read('word/document.xml');assert b'\\operatorname' not in xml and b'\\tag' not in xml and b'[REF' not in xml and b'[NEED' not in xml
 assert xml.count(b'<m:oMathPara>')==expected_equations
assert sha(src)==before
receipt=dict(status='DOCX_CREATED_RENDER_PENDING',line_spacing=2.0 if args.language=='en' else 1.5,source_md_sha256=before,docx_sha256=sha(out),native_display_equations=expected_equations,native_inline_math_objects=xml.count(b'<m:oMath>')-expected_equations,editable_tables=table_count,reference_items=len(re.split(r"## (?:References|参考文献)\n\n",t)[1].strip().split("\n\n")),planned_figure_legends=0,new_rendered_figures_in_docx=0,scientific_execution=0,builder_sha256=sha(Path(__file__)),source_md_unchanged=True)
args.receipt.write_text(json.dumps(receipt,ensure_ascii=False,indent=2)+'\n')
print(json.dumps(receipt,ensure_ascii=False))
