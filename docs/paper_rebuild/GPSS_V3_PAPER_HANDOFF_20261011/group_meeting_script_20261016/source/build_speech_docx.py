from pathlib import Path
import json,re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT=Path(__file__).resolve().parent.parent
data=json.loads((ROOT/'source/speech.json').read_text(encoding='utf-8'))
plan=json.loads((ROOT.parent/'组会汇报_20261015/source/SLIDE_PLAN.json').read_text(encoding='utf-8'))
assert len(data['slides'])==len(plan['slides'])==29
assert [s['title'] for s in data['slides']]==[s['title'] for s in plan['slides']]
OUT=ROOT/'output';OUT.mkdir(parents=True,exist_ok=True)

doc=Document()
section=doc.sections[0]
section.page_width=Cm(21.59);section.page_height=Cm(27.94)
section.top_margin=Cm(2.1);section.bottom_margin=Cm(2.0)
section.left_margin=Cm(2.3);section.right_margin=Cm(2.3)
section.header_distance=Cm(.8);section.footer_distance=Cm(.9)

def font(style,latin,east,size,bold=False,color='000000'):
    style.font.name=latin;style.font.size=Pt(size);style.font.bold=bold;style.font.color.rgb=RGBColor.from_string(color)
    rpr=style.element.get_or_add_rPr();rf=rpr.find(qn('w:rFonts'))
    if rf is None:rf=OxmlElement('w:rFonts');rpr.insert(0,rf)
    for k,v in [('ascii',latin),('hAnsi',latin),('eastAsia',east),('cs',latin)]:rf.set(qn('w:'+k),v)

font(doc.styles['Normal'],'Calibri','宋体',12)
normal=doc.styles['Normal'].paragraph_format
normal.line_spacing=1.2;normal.space_after=Pt(6);normal.widow_control=True
font(doc.styles['Title'],'Calibri','微软雅黑',23,True)
doc.styles['Title'].paragraph_format.space_after=Pt(8)
font(doc.styles['Subtitle'],'Calibri','微软雅黑',11,False,'555555')
font(doc.styles['Heading 1'],'Calibri','微软雅黑',15,True)
doc.styles['Heading 1'].paragraph_format.space_before=Pt(15)
doc.styles['Heading 1'].paragraph_format.space_after=Pt(6)
doc.styles['Heading 1'].paragraph_format.keep_with_next=True
font(doc.styles['Heading 2'],'Calibri','微软雅黑',12,True)
doc.styles['Heading 2'].paragraph_format.keep_with_next=True
doc.styles['Heading 2'].paragraph_format.space_before=Pt(10)
doc.styles['Heading 2'].paragraph_format.space_after=Pt(4)
for node in list(doc.styles.element.iter(qn('w:pBdr'))):
    node.getparent().remove(node)
for name,size in [('Cue',9.5),('Transition',11.5)]:
    style=doc.styles.add_style(name,1)
    style.base_style=doc.styles['Normal']
    font(style,'Calibri','微软雅黑' if name=='Cue' else '宋体',size,False,'666666' if name=='Cue' else '000000')
    style.paragraph_format.space_after=Pt(5)
    style.paragraph_format.keep_with_next=name=='Cue'

header=section.header.paragraphs[0]
header.text='紧凑四足机器人航向估计  组会汇报稿'
header.style=doc.styles['Cue']
for run in header.runs:run.font.color.rgb=RGBColor(0,0,0)
footer=section.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
run=footer.add_run();fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');run._r.addnext(fld)
doc.add_paragraph(data['title'],'Title')
doc.add_paragraph('对应29页组会PPT  '+data['date'],'Subtitle')
doc.add_paragraph('本稿围绕短基线方向与机身运动观测的条件互补展开，按现有PPT逐页编排。黑色正文可直接讲述，灰色指图提示供排练参考；每页末尾为衔接下一页的过渡句。主讲按约30分钟准备，备用答疑不计入主讲时间。')

md=['# '+data['title'],'',data['date']+' · 对应29页PPT · 主讲约30分钟，答疑另计','']
for item in data['slides']:
    clean=re.sub(r'[：:／/]',' ',item['title'])
    heading=f'第{item["page"]}页  {clean}'
    doc.add_paragraph(heading,'Heading 1')
    doc.add_paragraph('指图提示  '+item['cue'],'Cue')
    md.extend(['## '+heading,'','指图提示：'+item['cue'],''])
    for paragraph in item['paragraphs']:
        doc.add_paragraph(paragraph);md.extend([paragraph,''])
    if item['transition']:
        p=doc.add_paragraph(style='Transition')
        p.add_run('过渡  ').bold=True;p.add_run(item['transition'])
        md.extend(['过渡：'+item['transition'],''])

doc.add_page_break()
doc.add_paragraph('备用答疑','Heading 1')
doc.add_paragraph('以下内容在老师或同学提出相应问题时使用，不需要在主讲中逐条念出。','Cue')
md.extend(['# 备用答疑','','以下内容按现场问题选用，不计入主讲。',''])
for i,item in enumerate(data['qa'],1):
    title=f'{i}  {item["question"]}'
    doc.add_paragraph(title,'Heading 2')
    doc.add_paragraph(item['answer'])
    md.extend(['## '+title,'',item['answer'],''])

doc.core_properties.title=data['title']
doc.core_properties.subject='逐页组会口播与备用答疑'
doc.core_properties.author=''
doc.core_properties.keywords='GNSS INS 四足机器人 航向 组会讲稿'
doc.save(OUT/'LegSA-GINS_组会汇报稿_逐页口播版_20261016.docx')
(ROOT/'source/组会汇报稿_逐页正文.md').write_text('\n'.join(md),encoding='utf-8')
spoken=''.join(''.join(x['paragraphs'])+x['transition'] for x in data['slides'])
counts={'ppt_pages':29,'qa_questions':len(data['qa']),'spoken_nonspace_characters':len(re.sub(r'\s','',spoken)),
        'spoken_hanzi':len(re.findall(r'[\u4e00-\u9fff]',spoken)),'duration_estimate':'approximately 30 minutes including pointing and pauses',
        'source_pptx':'../组会汇报_20261015/output/pptx/LegSA-GINS_论文组会汇报_20261015.pptx'}
(ROOT/'source/SCRIPT_IDENTITY.json').write_text(json.dumps(counts,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(counts,ensure_ascii=False))
